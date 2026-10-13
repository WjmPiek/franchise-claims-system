"""External pay-over controls, opening balances and receipts; separate from premiums."""
import json
from datetime import datetime
from decimal import Decimal
from openpyxl import load_workbook
from central_calculations import normal
from direct_data import number

def ensure(store):
    key=(str(store.direct.path),id(store.direct.get_engine))
    if getattr(store,'_ledger_ready',None)==key:return
    with store.direct.connection() as conn:
        if store.direct.get_engine:store.direct.execute(conn,'SELECT pg_advisory_xact_lock(hashtext(:key))',{'key':'claims-ledger-schema-v1'})
        for sql in (
            'CREATE TABLE IF NOT EXISTS central_ledger_batches (id TEXT PRIMARY KEY,owner_id INTEGER NOT NULL,month TEXT NOT NULL,source_file TEXT NOT NULL,created_at TEXT NOT NULL)',
            'CREATE TABLE IF NOT EXISTS central_ledger_records (batch_id TEXT NOT NULL,franchise TEXT NOT NULL,record_json TEXT NOT NULL,PRIMARY KEY(batch_id,franchise))',
            'CREATE TABLE IF NOT EXISTS central_ledger_publications (month TEXT PRIMARY KEY,batch_id TEXT NOT NULL)'):
            store.direct.execute(conn,sql)
    store._ledger_ready=key

def read_ledger(path,sheet_name):
    book=load_workbook(path,read_only=True,data_only=True)
    formulas=load_workbook(path,read_only=True,data_only=False)
    try:
        if sheet_name not in book.sheetnames:raise ValueError('Choose the exact pay-over balance worksheet.')
        sheet=book[sheet_name];header=next(sheet.iter_rows(max_row=1,values_only=True));keys=[normal(v) for v in header]
        fields={'franchise':'Franchise','account':'Account Number / Reference Number','opening':'OUTSTANDING PAYOVER','payment':'BETALING','paid_on':'DATUM BETAAL','system_control':'System Pay-over'}
        if any(normal(v) not in keys for v in fields.values()):raise ValueError('The pay-over balance headings do not match the supplied example.')
        indices={k:keys.index(normal(v)) for k,v in fields.items()};seen=set()
        formula_rows=formulas[sheet_name].iter_rows(min_row=2,values_only=True)
        for n,(row,formula_row) in enumerate(zip(sheet.iter_rows(min_row=2,values_only=True),formula_rows),2):
            franchise=str(row[indices['franchise']] or '').strip()
            if not franchise:continue
            if not str(row[indices['account']] or '').strip():continue
            if franchise in seen:raise ValueError('Duplicate franchise balance rows; resolve before replacement.')
            seen.add(franchise)
            for key in ('opening','payment','system_control'):
                i=indices[key]
                if row[i] is None and str(formula_row[i] or '').startswith('='):raise ValueError(f'Balance source row {n}: {fields[key]} formula has no cached value. Recalculate and save the workbook first.')
            result={k:str(row[i] or '') for k,i in indices.items()}
            for key in ('opening','payment','system_control'):result[key]=str(number(row[indices[key]],fields[key],n))
            yield result
        if not seen:raise ValueError('No verified pay-over balance records found.')
    finally:book.close();formulas.close()

def publish(store,identifier,owner,month,filename,records):
    ensure(store)
    with store.direct.connection() as conn:
        if store.direct.get_engine:store.direct.execute(conn,'SELECT pg_advisory_xact_lock(hashtext(:key))',{'key':'ledger:'+month})
        store.reject_older_job(conn,identifier,month,'central_ledger_publications')
        if store.direct.execute(conn,'SELECT COUNT(*) FROM central_ledger_batches WHERE id=:id',{'id':identifier}).fetchone()[0]:return
        store.direct.execute(conn,'INSERT INTO central_ledger_batches VALUES (:id,:owner,:month,:source,:created)',dict(id=identifier,owner=owner,month=month,source=filename,created=datetime.utcnow().isoformat()))
        for record in records:store.direct.execute(conn,'INSERT INTO central_ledger_records VALUES (:batch,:franchise,:record)',dict(batch=identifier,franchise=record['franchise'],record=json.dumps(record)))
        store.direct.execute(conn,'INSERT INTO central_ledger_publications VALUES (:month,:batch) ON CONFLICT(month) DO UPDATE SET batch_id=excluded.batch_id',dict(month=month,batch=identifier))

def controls(store,month):
    ensure(store)
    with store.direct.connection() as conn:
        rows=store.direct.execute(conn,'SELECT r.record_json FROM central_ledger_records r JOIN central_ledger_publications p ON p.batch_id=r.batch_id WHERE p.month=:month',{'month':month})
        result={}
        for row in rows:
            value=json.loads(row[0]);key=normal(value['franchise'])
            if key in result:raise ValueError('Ambiguous franchise names in the balance source.')
            result[key]=value
        return result
