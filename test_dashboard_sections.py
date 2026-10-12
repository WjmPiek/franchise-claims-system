import io
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import pandas as pd
from flask import g
from openpyxl import load_workbook
from reportlab.pdfbase.pdfmetrics import stringWidth
os.environ['DATABASE_URL']=''
os.environ['APP_ENV']='local'
import app as claims
from report_files import table_workbook,workbook_pdf,workbook_preview


class DashboardSectionsTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.original=claims.LAST_RESULT
        cls.original_before=claims.app.before_request_funcs.copy()
        cls.engine_patch=patch.object(claims,'get_db_engine',return_value=None);cls.engine_patch.start()
        cls.relation_patch=patch.object(claims,'get_policy_relation_counts',return_value={'rows':[{'code':'MEM','label':'Main member','count':20}],'total':20});cls.relation_patch.start()
        def user():g.user={'id':999999,'role':'admin','is_active':True,'name':'Preview administrator','email':'preview@example.invalid'}
        claims.app.before_request_funcs={None:[user]}
        raw=pd.DataFrame([{'franchise':'Example Franchise','month':pd.Timestamp(f'{year}-{month:02d}-01'),'retail_premium':1000+month*10,'risk_premium':600+month*10,'original_risk_premium':650+month*10,'claims':100+month,'claim_count':2,'policy_qty':10+month,'r1_policy_fee_imported':10+month,'underwriter_2_1_fee':12} for year in (2025,2026) for month in (8,9,10)])
        raw['current_scenario']='100%'
        monthly,periods,portfolio=claims.analyse(raw,claims.DEFAULT_RATES,claims.DEFAULT_BOOK_VALUE)
        claims.LAST_RESULT={'raw':raw,'monthly':monthly,'periods':periods,'portfolio':portfolio,'rates':claims.DEFAULT_RATES,'book_rates':claims.DEFAULT_BOOK_VALUE}
        cls.client=claims.app.test_client()
    @classmethod
    def tearDownClass(cls):
        claims.LAST_RESULT=cls.original;claims.app.before_request_funcs=cls.original_before;cls.engine_patch.stop();cls.relation_patch.stop();cls.temp.cleanup()
    def test_dashboard_has_sidebar_links_and_no_import_report_comparison_or_map_panels(self):
        response=self.client.get('/dashboard');self.assertEqual(response.status_code,200,response.data[:200])
        for path in ('compare','top_franchise','imports','downloads','reports','graphs'):
            self.assertIn(('/overview/'+path).encode(),response.data)
        for removed in (b'id="importForm"',b'id="reportDownloadForm"',b'id="premiumParlourAnalysis"',b'id="topFranchises"',b'dashboardClientMiniMap',b'maps.googleapis.com'):
            self.assertNotIn(removed,response.data)
        self.assertIn(b'/policy-members?',response.data)
    def test_independent_pages_render_their_own_content(self):
        expected={'imports':b'id="importForm"','reports':b'View Book Value','compare':b'First date','top_franchise':b'id="topFranchises"','downloads':b'Saved Direct Reports','graphs':b'data-direct-series','claim_trend':b'Claim ratio','policy_trend':b'Policy / client quantity','relations':b'latest active member snapshot'}
        for section,text in expected.items():
            with self.subTest(section=section):
                response=self.client.get('/overview/'+section);self.assertEqual(response.status_code,200,response.data[:500]);self.assertTrue(text in response.data,section)
        self.assertIn(b'Second date',self.client.get('/overview/compare').data)
        self.assertIn(b'Difference',self.client.get('/overview/compare').data)
    def test_comparison_values_dates_and_missing_prior_data(self):
        source=claims.LAST_RESULT['monthly'];current=source[source.Month==pd.Timestamp('2026-10-01')]
        analysis=claims.build_premium_parlour_analysis(source,current,'Example Franchise')
        self.assertEqual((analysis['current_period'],analysis['last_year_period']),('Oct 2026','Oct 2025'))
        self.assertEqual(analysis['comparison_rows'][0]['difference'],0)
        missing=claims.build_premium_parlour_analysis(current,current,'Example Franchise')
        self.assertIsNone(missing['comparison_rows'][0]['previous']);self.assertIsNone(missing['comparison_rows'][0]['difference'])
    def test_report_catalog_financial_view_excel_and_pdf(self):
        for key in claims.FINANCIAL_REPORT_COLUMNS:
            args=f'?report_type={key}&report_period=month&report_month=2026-10'
            response=self.client.get('/export'+args+'&report_format=view');self.assertEqual(response.status_code,200)
            response=self.client.get('/export'+args);self.assertEqual(response.status_code,200)
            book=load_workbook(io.BytesIO(response.data),read_only=True);self.assertIn('Monthly Detail',book.sheetnames);book.close();response.close()
            response=self.client.get('/board_report'+args);self.assertEqual(response.status_code,200);self.assertTrue(response.data.startswith(b'%PDF'));response.close()
        for section in ('compare','top_franchise','claim_trend','policy_trend','relations'):
            for kind in ('xlsx','pdf'):
                response=self.client.get(f'/overview/{section}?download={kind}&premium_analysis_month=2026-10')
                self.assertEqual(response.status_code,200,section)
                self.assertTrue(response.data.startswith(b'%PDF' if kind=='pdf' else b'PK'),section);response.close()
    def test_complete_wide_workbook_pdf_and_pageable_preview(self):
        headers=['Franchise','Group','Policy','Surname','Full Names','ID','DOB','Relation','Risk','Retail','Product']
        values=[['Example','Martins',f'000{i:04d}','Malherbe','Member','0012312345678','02/03/1980','MEM',100.005,150.123,'Long family product description '*5] for i in range(61)]
        path=Path(self.temp.name)/'report.xlsx';path.write_bytes(table_workbook(headers,values).getvalue())
        self.assertEqual(len(workbook_preview(path,page=2)['rows']),11)
        pdf=workbook_pdf(path,'Direct Insurer Report');self.assertTrue(pdf.getvalue().startswith(b'%PDF'))
        proof=Path(__file__).parent.parent/'direct-data-check'/'report_layout_validation.pdf';proof.write_bytes(pdf.getvalue())

if __name__=='__main__':unittest.main()
