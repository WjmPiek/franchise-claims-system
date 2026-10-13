import ast
import os
import re
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from policy_layout import find_policy_sheet, policy_columns


class PolicyLayoutTests(unittest.TestCase):
    def setUp(self):
        names = {'read_policydata_streaming', 'looks_like_policydata_file',
                 'month_from_filename', 'display_source_filename', '_normalise_header',
                 'clean_money', 'parse_policy_transaction_sheet'}
        source = Path(__file__).resolve().parents[1].joinpath('app.py').read_text(encoding='utf-8')
        module = ast.Module(body=[node for node in ast.parse(source).body
                                if isinstance(node, ast.FunctionDef) and node.name in names],
                            type_ignores=[])
        self.env = dict(pd=pd, os=os, re=re, datetime=datetime, Decimal=Decimal,
                        ROUND_HALF_UP=ROUND_HALF_UP, load_workbook=load_workbook,
                        _clean_map_value=lambda value: value or '',
                        _clean_id_number=lambda value: value or '')
        exec(compile(module, 'policy-import-functions', 'exec'), self.env)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def workbook(self, header, rows, title_rows=0, cover=False):
        path = Path(self.temp.name) / '12345678-1234-1234-1234-123456789abc_PolicyData_20260901_to_20260930.xlsx'
        book = Workbook()
        sheet = book.active
        if cover:
            sheet.append(['Report information'])
            sheet = book.create_sheet('Transactions')
        for _ in range(title_rows):
            sheet.append(['Monthly receipt export'])
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        book.save(path)
        book.close()
        return path

    def test_aliases_whitespace_and_reordered_columns(self):
        header = [' MPIA ', 'Retail\nPremium', 'Member Relation', 'AUL Risk', 'Branch Name']
        rows = [[2, 150, 'MEM', 100, 'ERMELO'], [1, 30, 'EXT', 20, 'ERMELO']]
        path = self.workbook(header, rows, title_rows=2, cover=True)
        self.assertTrue(self.env['looks_like_policydata_file'](path))
        result = self.env['read_policydata_streaming'](path).iloc[0]
        self.assertEqual(result.franchise, 'ERMELO')
        self.assertEqual(result.month, pd.Timestamp(2026, 9, 1))
        self.assertEqual(result.retail_premium, 180)
        self.assertEqual(result.original_risk_premium, 120)
        self.assertAlmostEqual(result.risk_premium, 115.94)
        self.assertEqual(result.policy_qty, 2)
        self.assertEqual(self.env['LAST_POLICY_DETAIL_DF'].row_number.tolist(), [4, 5])
        fallback = self.env['parse_policy_transaction_sheet'](pd.DataFrame(rows, columns=header), str(path)).iloc[0]
        self.assertAlmostEqual(fallback.risk_premium, result.risk_premium)
        self.assertEqual(fallback.retail_premium, result.retail_premium)

    def test_risk_and_retail_premium_aliases_without_policydata_name(self):
        path = self.workbook(['Franchise', 'Relation', 'Risk Premium', 'Retail Premium', 'MPIA'],
                             [['ERMELO', 'MEM', 100, 150, 2]])
        renamed = path.with_name('Receipts_September_2026.xlsx')
        path.rename(renamed)
        self.assertTrue(self.env['looks_like_policydata_file'](renamed))
        self.assertEqual(self.env['read_policydata_streaming'](renamed).iloc[0].retail_premium, 150)

    def test_missing_required_column_is_actionable_and_not_guessed(self):
        path = self.workbook(['Franchise', 'Relation', 'Risk', 'Other Amount', 'MPIA'],
                             [['ERMELO', 'MEM', 100, 150, 2]])
        with self.assertRaisesRegex(ValueError, 'Required columns:.*Retail'):
            self.env['read_policydata_streaming'](path)
        unrelated = pd.DataFrame([['ERMELO'] + [100] * 22], columns=['Franchise'] + list(range(22)))
        self.assertTrue(self.env['parse_policy_transaction_sheet'](unrelated, str(path)).empty)

    def test_scan_is_bounded_and_empty_book_has_clear_error(self):
        path = self.workbook(['Franchise', 'Relation', 'Risk', 'Retail', 'MPIA'],
                             [['ERMELO', 'MEM', 100, 150, 2]], title_rows=25)
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            with self.assertRaisesRegex(ValueError, 'first 25 rows'):
                find_policy_sheet(book)
        finally:
            book.close()

    def test_legacy_xls_does_not_use_openpyxl(self):
        self.assertFalse(self.env['looks_like_policydata_file']('PolicyData_September_2026.xls'))


if __name__ == '__main__':
    unittest.main()
