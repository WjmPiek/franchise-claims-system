"""Consistent workbook previews and complete, paginated PDF exports."""
import io
from datetime import date,datetime
from decimal import Decimal
from openpyxl import load_workbook
import xlsxwriter
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import landscape,A4
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics

def display(value):
    if value is None:return ''
    if isinstance(value,(date,datetime)):return value.strftime('%d/%m/%Y')
    if isinstance(value,(float,Decimal)):
        whole,dot,fraction=format(Decimal(str(value)),',f').partition('.')
        return whole+'.'+fraction.rstrip('0').ljust(2,'0')
    return str(value)

def table_workbook(headers,rows,title='Report',extra_sheets=None):
    output=io.BytesIO()
    book=xlsxwriter.Workbook(output,{'in_memory':True,'strings_to_formulas':False,'strings_to_urls':False})
    head=book.add_format({'bold':True,'bg_color':'#E8EDF5'})
    for name,columns,values in [('Report',headers,rows),*(extra_sheets or [])]:
        sheet=book.add_worksheet(name[:31]);sheet.write_row(0,0,columns,head);sheet.freeze_panes(1,0);sheet.set_column(0,len(columns)-1,24)
        for n,row in enumerate(values,1):
            for col,value in enumerate(row):
                if value is None:continue
                if isinstance(value,Decimal):value=float(value)
                sheet.write(n,col,value)
        sheet.autofilter(0,0,len(values),len(columns)-1)
    book.close();output.seek(0)
    return output

def workbook_preview(path,sheet_name='',page=1):
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        if sheet_name and sheet_name not in book.sheetnames:raise ValueError('Choose an available worksheet.')
        sheet=book[sheet_name] if sheet_name else book[book.sheetnames[0]]
        count=max(0,(sheet.max_row or 1)-1);page=max(1,min(page,max(1,(count+49)//50)))
        headers=[display(v) for v in next(sheet.iter_rows(min_row=1,max_row=1,values_only=True))]
        rows=[[display(v) for v in row] for row in sheet.iter_rows(min_row=2+(page-1)*50,max_row=min(sheet.max_row or 1,1+page*50),values_only=True)] if count else []
        return {'sheets':book.sheetnames,'sheet':sheet.title,'headers':headers,'rows':rows,'count':count,'page':page}
    finally:book.close()

def workbook_pdf(path,title):
    """Render every row and column; repeat identifiers in panels for wide sheets."""
    output=io.BytesIO();pdf=canvas.Canvas(output,pagesize=landscape(A4),pageCompression=1);pdf.setTitle(title)
    width,height=landscape(A4);margin=32;usable=width-2*margin
    def wrap(value,space,size=8,bold=False):
        text=display(value).replace('\n',' ');font='Helvetica-Bold' if bold else 'Helvetica'
        lines=[];line=''
        for word in text.split():
            candidate=(line+' '+word).strip()
            if pdfmetrics.stringWidth(candidate,font,size)<=space:line=candidate;continue
            if line:lines.append(line);line=''
            for char in word:
                if line and pdfmetrics.stringWidth(line+char,font,size)>space:lines.append(line);line=''
                line+=char
        lines.append(line)
        return lines
    def footer():
        pdf.setFont('Helvetica',8);pdf.setFillColor(colors.HexColor('#596675'))
        pdf.drawString(margin,18,"Martin's Funerals - Claims");pdf.drawRightString(width-margin,18,str(pdf.getPageNumber()))
    def heading(label,sub=''):
        pdf.setFillColor(colors.HexColor('#18202A'));pdf.setFont('Helvetica-Bold',16)
        lines=wrap(label,usable,16,True)
        y=height-margin
        for line in lines:pdf.drawString(margin,y,line);y-=20
        if sub:
            pdf.setFont('Helvetica',9)
            for line in wrap(sub,usable,9):pdf.drawString(margin,y,line);y-=13
        return y-16
    heading("Martin's Funerals",'Claims reporting')
    pdf.setFont('Helvetica-Bold',22);y=height/2+40
    for line in wrap(title,usable,22,True):pdf.drawString(margin,y,line);y-=28
    pdf.setFont('Helvetica',10);pdf.drawString(margin,margin+42,'Prepared from the selected report data.');footer();pdf.showPage()
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        y=heading('Report summary','Rows below are counted per worksheet. Some worksheets contain overlapping data or summary rows.')
        for sheet in book:
            if y<margin+40:footer();pdf.showPage();y=heading('Report summary (continued)')
            pdf.setFont('Helvetica-Bold',10);pdf.drawString(margin,y,sheet.title)
            pdf.setFont('Helvetica',10);pdf.drawRightString(width-margin,y,f'{max(0,(sheet.max_row or 1)-1):,} worksheet rows');y-=22
            headers=list(next(sheet.iter_rows(min_row=1,max_row=1,values_only=True)))
            if 'Franchise' in headers and 'Policy Number' in headers:
                fields=[(name,headers.index(name)) for name in ('Risk','Retail','Fee') if name in headers]
                totals={name:Decimal(0) for name,col in fields}
                for row in sheet.iter_rows(min_row=2,values_only=True):
                    for name,col in fields:
                        value=row[col] if col<len(row) else None
                        if isinstance(value,(int,float,Decimal)):totals[name]+=Decimal(str(value))
                pdf.setFont('Helvetica',9)
                for line in wrap('  |  '.join(name+' (as exported): '+display(value) for name,value in totals.items()),usable,9):
                    if y<margin+30:footer();pdf.showPage();y=heading('Report summary (continued)')
                    pdf.drawString(margin,y,line);y-=15
                y-=8
        footer();pdf.showPage()
        for sheet in book:
            headers=list(next(sheet.iter_rows(min_row=1,max_row=1,values_only=True)))
            fixed=list(range(min(3,len(headers)))) if len(headers)>8 else []
            rest=list(range(len(fixed),len(headers)))
            groups=[fixed+rest[n:n+8-len(fixed)] for n in range(0,len(rest),8-len(fixed))] or [fixed]
            for panel,indices in enumerate(groups,1):
                cell_width=usable/max(1,len(indices))
                def start_table():
                    y=heading(sheet.title,f'{title} - columns panel {panel} of {len(groups)}')
                    cells=[wrap(headers[i],cell_width-10,8,True) for i in indices];hh=max([len(v) for v in cells] or [1])*10+12
                    pdf.setFillColor(colors.HexColor('#E8EDF5'));pdf.rect(margin,y-hh,usable,hh,fill=1,stroke=0);pdf.setFillColor(colors.HexColor('#18202A'));pdf.setFont('Helvetica-Bold',8)
                    for col,lines in enumerate(cells):
                        for line_no,line in enumerate(lines):pdf.drawString(margin+col*cell_width+5,y-12-line_no*10,line)
                    return y-hh
                y=start_table()
                for row in sheet.iter_rows(min_row=2,values_only=True):
                    cells=[wrap(row[i] if i<len(row) else '',cell_width-10) for i in indices];hh=max([len(v) for v in cells] or [1])*10+10
                    if y-hh<margin:footer();pdf.showPage();y=start_table()
                    pdf.setFillColor(colors.HexColor('#18202A'));pdf.setFont('Helvetica',8)
                    for col,lines in enumerate(cells):
                        for line_no,line in enumerate(lines):pdf.drawString(margin+col*cell_width+5,y-12-line_no*10,line)
                    pdf.setStrokeColor(colors.HexColor('#E1E5EF'));pdf.line(margin,y-hh,width-margin,y-hh);y-=hh
                footer();pdf.showPage()
    finally:book.close()
    pdf.save();output.seek(0);return output
