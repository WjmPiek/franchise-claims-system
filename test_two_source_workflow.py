import csv,json,sqlite3,unittest
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from contextlib import closing
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
import test_central_pages as pages
from test_direct_data import record
from direct_data import stage_source,prepare_outputs
from central_reports import catalog,generate,APPROVED_BRANCHES
from central_claims import verified_annual_sheet,ANNUAL_HEADERS

class TwoSourceWorkflow(pages.CentralPages):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.app.before_request_funcs[None].append(cls.app.extensions['approved_file_workflow'])
    def test_only_two_file_selectors_and_no_legacy_download_links(self):
        response=self.client.get('/overview/imports')
        self.assertEqual(response.status_code,200);self.assertEqual(response.data.count(b'type="file"'),2)
        self.assertIn(b'Annual Martins Funeral All Claims',response.data)
        response=self.client.get('/overview/downloads?date_from=2026-09-01&date_to=2026-09-30')
        self.assertIn(b'Monthly Pay-over per Franchise',response.data)
        for link in (b'/export?',b'/board_report?',b'/direct-data/reports',b'Complete ZIP',b'Converted_Source',b'Commission_Summary'):self.assertNotIn(link,response.data)
        admin=self.client.get('/overview/admin')
        self.assertNotIn(b'type="file"',admin.data);self.assertNotIn(b'Convert saved month',admin.data)
        self.assertIn(b'Save verified controls',admin.data)
    def test_old_import_and_export_routes_are_retired(self):
        self.assertEqual(self.client.get('/export').status_code,302)
        self.assertEqual(self.client.get('/direct-data/reports').status_code,302)
        self.assertEqual(self.client.post('/admin/payover/import-balances').status_code,410)
        self.assertEqual(self.client.post('/admin/direct/convert-saved').status_code,410)
        from report_workflow import RemoveOldLinks
        parser=RemoveOldLinks();parser.feed('<form action="/export"><button>Legacy</button></form><a href="/reports/central/x/0?format=pdf">PDF</a>')
        self.assertEqual(''.join(parser.output),'<a href="/reports/central/x/0?format=pdf">PDF</a>')
    def test_approved_upload_families_and_annual_actual_month_scope(self):
        with patch.object(self.queue,'create',return_value='test-job') as create:
            r=self.client.post('/imports/central',data={'kind':'claims','file':(BytesIO(b'fixture'),'Martins Funeral All Claims 1 Aug 2025 - 31 Aug 2026.xlsx')})
            self.assertEqual(r.status_code,202)
            args=create.call_args.args[2];self.assertEqual(args['replace_months'],[]);self.assertEqual(args['amount_header'],'Net amount paid via bank')
            r=self.client.post('/imports/central',data={'kind':'direct','file':(BytesIO(b'fixture'),'PolicyData_DTSTAMP_RECEIPTSFOR_20251001_to_20251031 (Example CSV Converted).xlsx')})
            self.assertEqual(r.status_code,202)
            r=self.client.post('/imports/central',data={'kind':'direct','file':(BytesIO(b'fixture'),'PolicyData_20260901_to_20260930_AUL.xlsx')})
            self.assertEqual(r.status_code,400)
    def test_embedded_tabs_comma_fragments_and_trailing_empty_fields(self):
        path=self.root/'embedded.xlsx';book=Workbook();sheet=book.active
        row=record(franchise='ERMELO',product='Gold Family Plan 31 to 55 Years R 40 000',risk='123.456',retail='200.005')
        row[14]='Road 1, Flat 2, Town';row[18:]=['','','','']
        fragments='\t'.join(map(str,row)).split(',');sheet.append(fragments);book.save(path)
        with closing(sqlite3.connect(':memory:')) as stage:
            control=stage_source(path,stage);source=json.loads(stage.execute('SELECT data FROM source_raw').fetchone()[0])
        self.assertEqual(control['rows'],1);self.assertEqual(control['risk'],Decimal('123.456'));self.assertEqual(control['retail'],Decimal('200.005'))
        self.assertEqual(source[14],'Road 1, Flat 2, Town');self.assertEqual(len(source),22);self.assertEqual(source[-1],'')
    def test_export_catalogue_contains_only_named_families_and_reference_columns(self):
        batch=self.store.publications()[0];items=catalog(self.store,batch)
        self.assertEqual({i['kind'] for i in items},{'payover','product_claims','policy','supplier','branch_calculation'})
        self.assertEqual(len([i for i in items if i['kind']=='payover']),1)
        self.assertEqual({i['destination'] for i in items if i['kind']=='branch_calculation'},set(APPROVED_BRANCHES))
        self.assertEqual({i['destination'] for i in items if i['kind']=='policy'},set(APPROVED_BRANCHES)|{'AUL'})
        for kind,columns in [('payover',17),('supplier',12),('branch_calculation',11)]:
            item=next(i for i in items if i['kind']==kind);path=self.root/(kind+'.xlsx');generate(self.store,batch,item,path)
            book=load_workbook(path,data_only=True)
            sheet=book['Monthly Pay-over'] if kind=='payover' else book['Sheet1'] if kind=='supplier' else book[item['destination']]
            self.assertEqual(sheet.max_column,columns)
            if kind=='payover':self.assertEqual(sheet['L1'].value,'CURRENT PAYOVER');self.assertEqual(sheet['Q1'].value,'TOTAL OUTSTANDING')
            if kind=='supplier':self.assertEqual(sheet['F1'].value,'Passport Nr');self.assertIsNone(sheet['F2'].value);self.assertEqual(sheet['J2'].value,13)
            if kind=='branch_calculation':self.assertEqual(sheet['A2'].number_format,'dd/mm/yyyy')
            book.close()
    def test_annual_claims_profile_requires_verified_demographic_fields(self):
        path=self.root/'annual.xlsx';book=Workbook();sheet=book.active;sheet.append(list(ANNUAL_HEADERS));book.save(path)
        self.assertEqual(verified_annual_sheet(path),'Sheet')
        sheet.delete_cols(1);book.save(path)
        with self.assertRaisesRegex(ValueError,'annual Martins'):verified_annual_sheet(path)
    def test_verified_legacy_historic_plans_and_orphan_known_plan2_rows(self):
        path=self.root/'known.csv'
        with path.open('w',newline='',encoding='utf-8') as f:csv.writer(f).writerows([record(product='PLAN E',cover='6000'),record(policy='ORPHAN',relation='EXT',product='Platinum Family Plan 31 to 55 Years R 50 000',cover='50000')])
        from central_validation import plan_assignments
        with closing(sqlite3.connect(':memory:')) as stage:
            stage_source(path,stage);plans,issues=plan_assignments(stage)
            self.assertEqual(plans,{1:'Historic',2:'Plan 2'})
            self.assertTrue(any('No MEM row' in note for note in issues[2]))
            stage.execute("UPDATE rows SET values_json=? WHERE n=1",(json.dumps(record(product='PLAN E',cover='50000')),))
            plans,issues=plan_assignments(stage)
            self.assertTrue(any(note.startswith('[BLOCK] Legacy') for note in issues[1]))
    def test_manual_controls_keep_other_franchises_and_do_not_create_a_file_import(self):
        response=self.client.post('/admin/payover/save-controls',data={'month':'2026-09','franchise':'MARTINS FUNERALS ERMELO','account':'TEST-REF','opening':'20','payment':'10','paid_on':'2026-09-20'})
        self.assertEqual(response.status_code,302)
        from central_ledger import controls
        data=controls(self.store,'2026-09')['MARTINSFUNERALSERMELO'];self.assertEqual(data['opening'],'20');self.assertIsNone(data['system_control'])

if __name__=='__main__':unittest.main()
