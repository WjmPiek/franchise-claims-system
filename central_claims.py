"""Paid transactions identified by headers, never by Excel column letters."""
import json,re,uuid
from datetime import datetime,date,timedelta
from decimal import Decimal,InvalidOperation
from pathlib import Path
from openpyxl import load_workbook
from central_calculations import normal

PAID_DATES={'DATEOFCLAIMPAID','CLAIMPAIDDATE','PAIDDATE'}
BRANCHES={'PARTICIPATINGGROUP','FRANCHISE','FRANCHISENAME','BRANCH'}
VERIFIED_AMOUNT_HEADER='Net amount paid via bank'
VERIFIED_NET_OFF_HEADER='Amount net off against claim'
VERIFIED_BANK_HEADER='Net amount paid via bank'

def component(source,heading):
    values=[v for k,v in source.items() if normal(k)==normal(heading)]
    if len(values)>1:raise ValueError('Duplicate claims component heading: '+heading)
    value=values[0] if values else ''
    if value in ('',None,'-'):return Decimal(0)
    text=str(value).strip().replace('R','').replace(' ','').replace(',','')
    if text.startswith('(') and text.endswith(')'):text='-'+text[1:-1]
    try:
        amount=Decimal(text)
        if not amount.is_finite():raise InvalidOperation
        return amount
    except InvalidOperation:raise ValueError('Invalid claims component: '+heading) from None

def columns(values,amount_header):
    keys=[normal(v) for v in values]
    amount=normal(amount_header)
    if not amount:return None
    if keys.count(amount)>1:raise ValueError('The selected paid amount heading appears more than once.')
    dates=[i for i,k in enumerate(keys) if k in PAID_DATES];branches=[i for i,k in enumerate(keys) if k in BRANCHES]
    if not dates or not branches or amount not in keys:return None
    if len(dates)>1 or len(branches)>1:raise ValueError('Multiple paid-date/franchise fields; specify the transaction worksheet.')
    return {'franchise':branches[0],'date':dates[0],'amount':keys.index(amount),'product':next((i for i,k in enumerate(keys) if k in ('PRODUCT','PRODUCTNAME','PLANNAME')),None)}

def paid_date(value,row_number):
    if isinstance(value,(datetime,date)):return value
    if isinstance(value,(float,int)):return datetime(1899,12,30)+timedelta(days=value)
    text=str(value or '').strip()
    for fmt in ('%Y-%m-%d','%d/%m/%Y','%d-%m-%Y','%Y/%m/%d'):
        try:return datetime.strptime(text.split(' ')[0],fmt)
        except ValueError:pass
    raise ValueError(f'Claims source row {row_number}: invalid Date Of Claim Paid; previous months are preserved.')

def source_paid_on(source):
    value=next((v for k,v in source.items() if normal(k) in PAID_DATES),'')
    return paid_date(value,0).strftime('%Y-%m-%d') if value else ''

def read_transactions(path,amount_header,sheet_name='',controls=None,event_sink=None):
    if not amount_header.strip():raise ValueError('Specify the verified heading containing the amount actually paid.')
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        matches=[]
        for sheet in book:
            if sheet_name and sheet.title!=sheet_name:continue
            for n,row in enumerate(sheet.iter_rows(max_row=25,values_only=True),1):
                found=columns(row,amount_header)
                if found:matches.append((sheet,n,row,found));break
        if len(matches)!=1:raise ValueError('Select one transaction worksheet containing the paid-date, franchise and verified paid amount headings. Pivot totals are not transaction sources.')
        sheet,header_n,header,index=matches[0]
        status_index=next((i for i,key in enumerate(header) if normal(key)=='CLAIMSTATUS'),None)
        for n,row in enumerate(sheet.iter_rows(min_row=header_n+1,values_only=True),header_n+1):
            if not any(v not in (None,'') for v in row):continue
            if controls is not None:controls['source_rows']=controls.get('source_rows',0)+1
            source={str(key or f'Column {i+1}'):str(row[i] if i<len(row) and row[i] is not None else '') for i,key in enumerate(header)}
            if event_sink is not None:event_sink(n,source)
            amount=component(source,amount_header)
            netoff=component(source,VERIFIED_NET_OFF_HEADER)
            status=str(row[status_index] or '').casefold() if status_index is not None else ''
            if amount==0:
                if netoff and row[index['date']] not in (None,''):
                    if controls is not None:controls['netoff_only_rows']=controls.get('netoff_only_rows',0)+1
                else:
                    if controls is not None:
                        controls['nonfinancial_rows']=controls.get('nonfinancial_rows',0)+1
                        if component(source,'Final Amount'):
                            controls['unpaid_rows']=controls.get('unpaid_rows',0)+1
                            controls['unpaid_final_amount']=str(Decimal(controls.get('unpaid_final_amount','0'))+component(source,'Final Amount'))
                        if netoff:
                            controls['undated_netoff_rows']=controls.get('undated_netoff_rows',0)+1
                            controls['undated_netoff_amount']=str(Decimal(controls.get('undated_netoff_amount','0'))+netoff)
                            controls.setdefault('undated_netoff_source_rows',[]).append(n)
                    continue
            if row[index['date']] in (None,'') and any(marker in status for marker in ('repudiated','awaiting outstanding information')):
                if component(source,VERIFIED_NET_OFF_HEADER) or component(source,VERIFIED_BANK_HEADER):
                    raise ValueError(f'Claims source row {n}: payment components exist without a paid date; resolve before publication.')
                if controls is not None:
                    controls['unpaid_rows']=controls.get('unpaid_rows',0)+1
                    controls['unpaid_final_amount']=str(Decimal(controls.get('unpaid_final_amount','0'))+amount)
                continue
            franchise=str(row[index['franchise']] or '').strip()
            if not franchise or normal(franchise) in ('TOTAL','GRANDTOTAL'):
                raise ValueError(f'Claims source row {n}: missing franchise or summary total in transaction worksheet.')
            month=paid_date(row[index['date']],n).strftime('%Y-%m')
            product=str(row[index['product']] or '').strip() if index['product'] is not None else ''
            yield {'n':n,'month':month,'franchise':franchise,'product':product,'amount':str(amount),
              'source':source}
    finally:book.close()

def publish_claims(store,stage,source_file,owner,amount_header,replace_months,batch_id=None):
    """Validate and stage first; replace the expressly selected paid months atomically."""
    store.ensure();batch=batch_id or uuid.uuid4().hex
    actual={row[0] for row in stage.execute('SELECT DISTINCT month FROM claims')}
    if not actual:raise ValueError('No nonzero paid transactions found; existing claims months are preserved.')
    scope=set(replace_months)
    if not scope or not scope.issubset(actual):raise ValueError('Choose complete paid months present in this file. Confirm a month with no financial claim records separately in Admin. Other source months remain archived without replacing their published data.')
    with store.direct.connection() as conn:
        exists=store.direct.execute(conn,'SELECT COUNT(*) FROM central_claim_batches WHERE id=:id',{'id':batch}).fetchone()[0]
        if exists:return {'batch_id':batch,'months':sorted(scope),'rows':stage.execute('SELECT COUNT(*) FROM claims').fetchone()[0],'recovered':True}
        if store.direct.get_engine:
            for month in sorted(scope):store.direct.execute(conn,'SELECT pg_advisory_xact_lock(hashtext(:key))',{'key':'claims:'+month})
        for month in sorted(scope):store.reject_older_job(conn,batch,month,'central_claim_publications')
        store.direct.execute(conn,'INSERT INTO central_claim_batches VALUES (:id,:source,:owner,:created,:amount,:status)',dict(id=batch,source=source_file,owner=owner,created=datetime.utcnow().isoformat(),amount=amount_header,status='Published'))
        for n,month,franchise,product,amount,source,issues in stage.execute('SELECT * FROM claims ORDER BY n'):
            store.direct.execute(conn,'INSERT INTO central_claims VALUES (:batch,:n,:month,:franchise,:product,:amount,:source,:issues)',dict(batch=batch,n=n,month=month,franchise=franchise,product=product,amount=amount,source=source,issues=issues))
            fields=json.loads(source)
            store.direct.execute(conn,'INSERT INTO central_claim_components VALUES (:batch,:n,:netoff,:bank)',dict(batch=batch,n=n,netoff=str(component(fields,VERIFIED_NET_OFF_HEADER)),bank=str(component(fields,VERIFIED_BANK_HEADER))))
        if stage.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='claim_events'").fetchone():
            from claim_insights import persist_events
            persist_events(store,conn,batch,stage.execute('SELECT n,source FROM claim_events ORDER BY n'))
        for month in scope:
            store.direct.execute(conn,'INSERT INTO central_claim_publications VALUES (:month,:batch) ON CONFLICT(month) DO UPDATE SET batch_id=excluded.batch_id',dict(month=month,batch=batch))
    return {'batch_id':batch,'months':sorted(scope),'unpublished_source_months':sorted(actual-scope),'rows':stage.execute('SELECT COUNT(*) FROM claims').fetchone()[0]}
