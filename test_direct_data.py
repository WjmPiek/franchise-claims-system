import csv
import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from decimal import Decimal
from flask import Flask, g
from openpyxl import load_workbook
from direct_data import HEADERS, VOUCHER_PRODUCT, VOUCHER_HEADERS, allocation, row_fees, stage_source, prepare_outputs, register_direct_data


def record(franchise='MARTINS FUNERALS ERMELO', policy='000011', relation='MEM', product='Platinum Family Plan R 50 000', risk='100.005', retail='150.123', mpia='1', cover='50000', phone='0720001234'):
    return [franchise, 'Martins Funerals', policy, 'Malherbé', 'Example Member', '0012312345678',
            '1980-03-02 00:00:00.000', '', relation, '2020-01-03 00:00:00.000', cover, product,
            risk, retail, 'Example address', phone, '0012345', mpia, '', '', '0001', 'Agent']


class DirectDataTest(unittest.TestCase):
    def test_faithnet_thresholds_and_mem_only(self):
        for mpia, expected in [('1',40),('2',40),('3',80),('5',80)]:
            r = record(franchise='FAITHNETSA', product='1 + 9 (Member and 9 nominated Family members)',risk='300',mpia=mpia)
            self.assertEqual(row_fees(r,'over_two_double','mpia')[0], expected)
        self.assertEqual(row_fees(record(franchise='FAITHNETSA',product='Single Plan',mpia='3'),'over_two_double','mpia')[0],20)
        r=record(franchise='FAITHNETSA',relation='EXT',risk='-20',mpia='3')
        self.assertEqual(row_fees(r,'over_two_double','mpia')[:4],(0,0,0,Decimal('-20')))
        r=record(franchise='FAITHNETSA',product='1 + 5',risk='200',mpia='2')
        self.assertEqual(row_fees(r,'over_two_double','mpia')[3],Decimal('154.68'))
        self.assertEqual(row_fees(record(risk='0'),'over_two_double','mpia')[3],0)

    def test_all_destinations_and_pretoria_aliases(self):
        for name in ['PRETORIA-NORTH','Pretoria Noord','HAMMANSKRAAL','SOSHANGUVE']:
            self.assertEqual(allocation('MARTINS FUNERALS '+name),'Pretoria North Group')
        self.assertEqual(allocation('MARTINS FUNERALS CARLETONVILLE & KEMPTON PARK'),'Carletonville')
        self.assertEqual(allocation('FaithnetSA'),'AUL')
        self.assertEqual(allocation('ERMELO'),'AUL')

    def test_real_csv_layout_precision_policy_split_privacy_and_types(self):
        rows = [record(),record(relation='DEP',product='Extra Child Cover',risk='0',retail='0',cover='2500'),
                record(relation='MEMBEN',product=VOUCHER_PRODUCT,risk='17',retail='20',mpia='2',phone=''),
                record(franchise='MARTINS FUNERALS ALBERTON',product='Single Plan',risk='20',retail='30',cover='10000',phone='0810000000')]
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'PolicyData_20251101_to_20251130.csv'
            with source.open('w',encoding='cp850',newline='') as f:
                csv.writer(f,delimiter='\t').writerows(rows)
            with sqlite3.connect(':memory:') as conn:
                control=stage_source(source,conn)
                self.assertEqual(control,{'rows':4,'risk':Decimal('137.005'),'retail':Decimal('200.123')})
                files,archive,totals,issues=prepare_outputs(conn,root/'reports',source.name)
            self.assertEqual(sum(v[0] for v in totals.values()),4)
            self.assertEqual(sum(v[1] for v in totals.values()),control['risk'])
            self.assertEqual(sum(v[2] for v in totals.values()),control['retail'])
            self.assertFalse(issues)
            self.assertTrue(archive.exists())
            wb=load_workbook(files['Insurer Sheets'],data_only=True)
            self.assertEqual(wb.sheetnames[0],'AUL Totals')
            self.assertIn('Pretoria North Group',wb.sheetnames)
            self.assertEqual(wb['AUL All Data'].max_row,3)
            self.assertEqual(wb['Alberton'].max_row,2)
            for sheet in wb:
                self.assertFalse(any(VOUCHER_PRODUCT in r for r in sheet.iter_rows(values_only=True)))
            wb.close()
            self.assertIn('Martins Funeral Plan 2 (New)',files['AUL Plan 2 New'].name)
            wb=load_workbook(files['Converted Source'],data_only=True)
            ws=wb.active
            self.assertEqual(ws['F2'].value,'0012312345678')
            self.assertEqual(ws['P2'].value,'0720001234')
            self.assertEqual(ws['D2'].value,'Malherbé')
            self.assertEqual(ws['G2'].number_format,'dd/mm/yyyy')
            self.assertEqual(ws['M2'].value,100.005)
            self.assertNotIn('R',ws['M2'].number_format)
            wb.close()
            wb=load_workbook(files['Vouchers'],data_only=True)
            self.assertEqual([c.value for c in wb.active[1]],VOUCHER_HEADERS)
            self.assertEqual(wb.active['I2'].value,13)
            self.assertEqual(wb.active['J2'].value,'0720001234')
            self.assertNotIn('ID Number',[c.value for c in wb.active[1]])
            wb.close()
            wb=load_workbook(files['AUL All Data'],data_only=True)
            self.assertEqual(wb.active.max_row,3)
            self.assertNotIn(VOUCHER_PRODUCT,[r[11] for r in wb.active.iter_rows(min_row=2,values_only=True)])
            wb.close()
            wb=load_workbook(files['AUL Plan 2 New'],data_only=True)
            self.assertEqual(wb.active.max_row,3)  # DEP stays with its MEM plan despite Cover 2500.
            wb.close()
            wb=load_workbook(files['AUL Totals'],data_only=True)
            self.assertEqual(wb.active['B2'].value,2)
            self.assertEqual(wb.active['B3'].value,2)
            wb.close()

    def test_missing_mem_and_ambiguous_phone_are_visible(self):
        rows=[record(),record(phone='0730000000'),record(relation='MEMBEN',product=VOUCHER_PRODUCT,risk='8.50',retail='10'),
              record(policy='ORPHAN',relation='EXT',product='Platinum Family R 50 000')]
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'source.csv'
            with source.open('w',encoding='utf-8',newline='') as f:csv.writer(f).writerows([HEADERS,*rows])
            with sqlite3.connect(':memory:') as conn:
                stage_source(source,conn)
                files,archive,totals,issues=prepare_outputs(conn,root/'out',source.name)
            self.assertEqual(len(issues),2)
            wb=load_workbook(files['Vouchers'],data_only=True)
            self.assertIsNone(wb.active['J2'].value)
            wb.close()

    def test_route_requires_admin_and_download_belongs_to_creator(self):
        with tempfile.TemporaryDirectory() as temp:
            app=Flask(__name__,template_folder=str(Path(__file__).parent/'templates'));app.secret_key='test'
            who={'id':1,'role':'admin'}
            @app.before_request
            def user():g.user=who
            app.add_url_rule('/dashboard','dashboard',lambda:'dashboard')
            app.add_url_rule('/api/import/session','import_session',lambda:'session')
            register_direct_data(app,temp)
            client=app.test_client()
            self.assertIn(b'Direct Data Preparation',client.get('/direct-data').data)
            content=io.StringIO();csv.writer(content,delimiter='\t').writerow(record())
            response=client.post('/direct-data',data={'data_month':'2026-09','direct_file':(io.BytesIO(content.getvalue().encode()),'source.csv')})
            self.assertEqual(response.status_code,200,response.data)
            result=client.post('/direct-data/reports/'+response.json['job']).json
            url=f"/direct-data/download/{result['job']}/{result['files'][0]}"
            download=client.get(url);self.assertEqual(download.status_code,200);download.close()
            stream=client.post('/direct-data',headers={'X-Import-Progress':'stream'},
                               data={'data_month':'2026-09','direct_file':(io.BytesIO(content.getvalue().encode()),'source.csv')})
            self.assertEqual(stream.mimetype,'application/x-ndjson')
            updates=[json.loads(line) for line in stream.data.splitlines()]
            self.assertTrue(updates[-1]['ok'])
            self.assertEqual(updates[-1]['source']['rows'],1)
            self.assertEqual(updates[-1]['files'],[])
            self.assertIn('job',updates[-1])
            who['id']=2;self.assertEqual(client.get(url).status_code,404)
            who['role']='viewer';self.assertEqual(client.get('/direct-data').status_code,403)

if __name__=='__main__':unittest.main()
