import csv
import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from flask import Flask,g
from openpyxl import load_workbook
from direct_data import stage_source,prepare_outputs,register_direct_data
from direct_analytics import DirectStore
from test_direct_data import record


class SavedDirectReportsTest(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def stage(self,rows):
        path=self.root/'PolicyData_20251001_to_20251031.csv'
        with path.open('w',newline='',encoding='utf-8') as handle:csv.writer(handle).writerows(rows)
        conn=sqlite3.connect(':memory:');stage_source(path,conn)
        return conn
    def test_separate_products_sharing_policy_number_follow_their_own_member(self):
        family='Gold Family Plan 31 to 55 Years R 40 000';cremation='Martin`s Cremation Plan For Families'
        conn=self.stage([record(policy='GS16432',product=family,cover='40000'),record(policy='GS16432',product=cremation,cover='20000'),record(policy='GS16432',relation='SPS',product=family,cover='10000',risk='0')])
        files,archive,totals,issues=prepare_outputs(conn,self.root/'reports','October.csv')
        plans=conn.execute('SELECT n,plan FROM prepared ORDER BY n').fetchall()
        self.assertEqual(plans,[(1,'Plan 2'),(2,'Historic'),(3,'Plan 2')])
        self.assertEqual(totals['AUL'][0],3)
        for name,count in [('AUL Plan 2 New',3),('AUL Historic',2)]:
            book=load_workbook(files[name],read_only=True);self.assertEqual(book.active.max_row,count);book.close()
        conn.close()
    def test_same_product_contradiction_and_ambiguous_dependent_still_require_review(self):
        for rows in [[record(product='Single Plan',cover='40000'),record(product='Single Plan',cover='20000')],[record(product='Family R 40 000',cover='40000'),record(product='Cremation',cover='20000'),record(relation='EXT',product='Extra member',cover='5000')]]:
            conn=self.stage(rows)
            with self.assertRaisesRegex(ValueError,'000011'):prepare_outputs(conn,self.root/'reports','October.csv')
            conn.close()
    def test_download_recreation_uses_stored_amounts_and_does_not_recalculate(self):
        conn=self.stage([record(franchise='FAITHNETSA',product='1 + 9',risk='300',mpia='3'),record(relation='MEMBEN',product='GROCERY DISCOUNT COUPONS FOR MARTINS CLIENTS',risk='17',mpia='2')])
        original=prepare_outputs(conn,self.root/'original','October.csv')
        store=DirectStore(self.root/'analytics.db');store.save(conn,'saved-batch',1,'2025-10','October.csv');conn.close()
        restored=sqlite3.connect(':memory:');batch=store.restore_stage(1,'saved-batch',restored)
        with patch('direct_data.row_fees',side_effect=AssertionError('Must use stored calculations')):
            recreated=prepare_outputs(restored,self.root/'recreated',batch['source_file'],saved=True)
        self.assertEqual(original[2],recreated[2]);self.assertEqual(original[3],recreated[3])
        for name in original[0]:
            before=load_workbook(original[0][name],read_only=True,data_only=True);after=load_workbook(recreated[0][name],read_only=True,data_only=True)
            self.assertEqual(before.sheetnames,after.sheetnames)
            for title in before.sheetnames:self.assertEqual(list(before[title].values),list(after[title].values))
            before.close();after.close()
        restored.close()
    def test_history_recreate_and_owner_access(self):
        app=Flask(__name__,template_folder=str(Path(__file__).parent/'templates'));app.secret_key='tests'
        who={'id':1,'role':'admin'}
        @app.before_request
        def user():g.user=who
        app.add_url_rule('/dashboard','dashboard',lambda:'dashboard')
        app.add_url_rule('/api/import/session','import_session',lambda:'session')
        register_direct_data(app,self.root)
        client=app.test_client();source=io.StringIO();csv.writer(source).writerow(record())
        response=client.post('/direct-data',data={'data_month':'2025-10','direct_file':(io.BytesIO(source.getvalue().encode()),'October.csv')})
        self.assertEqual(response.status_code,200,response.data);batch=response.json['job']
        self.assertIn(b'Recreate 2025-10 downloads',client.get('/direct-data/reports').data)
        result=client.post('/direct-data/reports/'+batch)
        self.assertEqual(result.status_code,200,result.data)
        self.assertIn('Direct_reports.zip',result.json['files'])
        downloaded=client.get('/direct-data/download/'+result.json['job']+'/Direct_reports.zip')
        self.assertEqual(downloaded.status_code,200);downloaded.close()
        self.assertIn(b'category=r1',client.get('/direct-data/analytics?category=r1').data)
        who['id']=2;self.assertNotIn(b'October.csv',client.get('/direct-data/reports').data)
        self.assertEqual(client.post('/direct-data/reports/'+batch).status_code,400)
        who['role']='viewer';self.assertEqual(client.get('/direct-data/reports').status_code,403)

if __name__=='__main__':unittest.main()
