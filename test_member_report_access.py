import io
import unittest
from flask import Flask,g
import pandas as pd
from sqlalchemy import create_engine,text
from openpyxl import load_workbook
from current_members import register_current_members
from policy_reports import register_policy_reports

class MemberReportAccessTest(unittest.TestCase):
    def test_current_relation_drilldown_paginates_and_enforces_scope(self):
        engine=create_engine('sqlite:///:memory:')
        with engine.begin() as conn:
            conn.execute(text('CREATE TABLE app_policydata_current_members (franchise_name TEXT,relation TEXT,member_name TEXT,policy_number TEXT,id_number TEXT,age INTEGER,age_band TEXT,import_month TEXT,active BOOLEAN)'))
            conn.execute(text('INSERT INTO app_policydata_current_members VALUES (:f,:r,:name,:policy,:id,30,:band,:month,true)'),[{'f':branch,'r':'MEM','name':'Member '+str(i),'policy':f'000{i:04d}','id':'0012312345678','band':'30','month':'2026-10-01'} for branch in ('A','B') for i in range(61)])
        app=Flask(__name__);user={'id':1,'role':'viewer'}
        @app.before_request
        def who():g.user=user
        app.add_url_rule('/overview/<section>','overview_page',lambda section:'overview')
        register_current_members(app,lambda:engine,text,lambda frame:frame[frame.Franchise=='A'])
        client=app.test_client();response=client.get('/policy-members?relation=MEM&page=2')
        self.assertEqual(response.status_code,200);self.assertIn(b'61 matching active members',response.data);self.assertIn(b'0000060',response.data)
        self.assertEqual(client.get('/policy-members?franchise=B').status_code,400)
        self.assertNotIn(b'<td>B</td>',response.data)
        self.assertIn(b'1 matching active members',client.get('/policy-members?q=0000060').data)
        user.clear();self.assertEqual(client.get('/policy-members').status_code,403);engine.dispose()
    def test_relation_excel_pdf_preserve_totals_and_external_controls(self):
        frame=pd.DataFrame([{'import_month':pd.Timestamp('2026-10-01'),'franchise':f,'relation':'MEM','original_risk_premium':100.005,'retail_premium':150.123,'r1_policy_fee':1,'adv_fund_2_1_fee':2.08,'new_risk_premium':96.93} for f in ('A','B')])
        app=Flask(__name__);app.add_url_rule('/dashboard','dashboard',lambda:'home')
        register_policy_reports(app,lambda:None,text,lambda:frame,lambda rows:rows[rows.Franchise=='A'])
        client=app.test_client();args='?month=2026-10&external_risk=100.01&external_retail=150.12'
        response=client.get('/policy_reports'+args+'&download=xlsx');self.assertEqual(response.status_code,200)
        book=load_workbook(io.BytesIO(response.data),read_only=True,data_only=True);self.assertIn('Reconciliation',book.sheetnames)
        rows=list(book['Report'].values);self.assertEqual(rows[1][1],'All');self.assertEqual(rows[1][4],100.005);self.assertEqual(rows[2][3],1);book.close()
        response=client.get('/policy_reports'+args+'&download=pdf');self.assertEqual(response.status_code,200);self.assertTrue(response.data.startswith(b'%PDF'))
        self.assertEqual(client.get('/policy_reports?franchise=B').status_code,400)

if __name__=='__main__':unittest.main()
