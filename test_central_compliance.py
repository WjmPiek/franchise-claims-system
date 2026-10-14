import csv,io,json,sqlite3,tempfile,unittest
from pathlib import Path
from decimal import Decimal
from unittest.mock import patch
from contextlib import closing
from flask import Flask,g
from openpyxl import Workbook,load_workbook
from direct_data import register_direct_data,stage_source,prepare_outputs
from direct_analytics import DirectStore
from central_calculations import effective_mpia,fees,ratio,retail_rank,VERSION
from central_store import CentralStore
from central_reports import catalog_all as catalog,generate
from central_claims import read_transactions,publish_claims
from central_jobs import JobQueue,register_jobs
from central_routes import register_central
from import_audit import ImportAudit
from test_direct_data import record

class CentralComplianceTest(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.direct=DirectStore(self.root/'data.sqlite');self.store=CentralStore(self.direct);self.direct.central=self.store
    def tearDown(self):self.temp.cleanup()
    def import_rows(self,rows,batch='batch1',month='2026-09'):
        path=self.root/(batch+'.csv')
        with path.open('w',encoding='utf-8',newline='') as f:csv.writer(f).writerows(rows)
        with closing(sqlite3.connect(self.root/(batch+'.sqlite'))) as stage:
            controls=stage_source(path,stage)
            prepare_outputs(stage,self.root/(batch+'-unused'),path.name,'mpia','mpia','policy',generate_files=False)
            self.direct.save(stage,batch,1,month,path.name)
        return controls
    def test_confirmed_rules_and_zero_basis(self):
        for source in ('',0,-2,'NULL'):self.assertEqual(effective_mpia(source),1)
        for mpia,amount in [('1',40),('2',80),('3',120),('5',200)]:
            row=record(franchise='FAITHNETSA',product='1 + 5',risk='300',mpia=mpia)
            self.assertEqual(fees(row)[0],amount)
        row=record(franchise='FAITHNETSA',product='1 + 5',risk='200',retail='300',mpia='2')
        self.assertEqual(fees(row),(Decimal(80),Decimal(2),Decimal('2.48'),Decimal('115.52'),Decimal(0),Decimal(0)))
        self.assertIsNone(ratio(100,0)['value']);self.assertIn('Undefined',ratio(100,0)['status']);self.assertIn('Not applicable',ratio(0,0)['status'])
        self.assertEqual([ratio(x,100)['colour'] for x in (49,50,69,70)],['blue','green','green','red'])
    def test_shared_identifiers_preserved_and_replacement_archived(self):
        rows=[record(franchise='FAITHNETSA',product='1 + 5',risk='200',retail='300',mpia='2')]*2
        self.import_rows(rows)
        summary=self.store.summaries()[0]['summary'];self.assertEqual(summary['rows'],2)
        self.assertEqual(Decimal(summary['faithnet']),160);self.assertEqual(Decimal(summary['paid_commissions']),Decimal('98.96'))
        self.assertEqual(Decimal(summary['admin']),270)
        with patch('central_calculations.fees',side_effect=AssertionError('Reading must never recalculate fees')):
            records=list(self.store.records('batch1'));self.assertEqual(len(records),2)
        self.import_rows([record(risk='400')],'replacement')
        self.assertEqual(len(self.store.publications()),1)
        self.assertEqual(self.store.publications()[0]['id'],'replacement')
        with self.direct.connection() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM central_records WHERE batch_id=?',('batch1',)).fetchone()[0],2)
    def test_source_mpia_originals_all_relations_and_conservation(self):
        rows=[record(mpia=''),record(relation='DEP',risk='12.345',retail='20.005'),record(relation='',risk='-2.333',retail='-1.2'),record(relation='MEMBEN',product='GROCERY DISCOUNT COUPONS FOR MARTINS CLIENTS',risk='17',retail='20',mpia='2')]
        control=self.import_rows(rows)
        records=list(self.store.records('batch1'));self.assertEqual(records[0]['calculation']['source_mpia'],'')
        self.assertEqual(records[0]['calculation']['calculation_mpia'],'1')
        v=self.store.summaries()[0]['summary']
        self.assertEqual(Decimal(v['original_risk']),control['risk']);self.assertEqual(Decimal(v['source_retail']),control['retail'])
        self.assertEqual(Decimal(v['r1']),1);self.assertEqual(Decimal(v['policy_qty']),1)
        self.assertEqual(Decimal(v['supplier']),13);self.assertEqual(Decimal(v['retained']),4)
        self.assertEqual(Decimal(v['percentage_commission']),Decimal(v['retail'])*Decimal('.15'))
        self.assertEqual(Decimal(v['original_risk']),sum((Decimal(v[k]) for k in ('payover','faithnet','r1','adv','supplier','retained','reconciliation_adjustment')),Decimal(0)))
    def test_split_unknown_product_and_individual_export_schemas(self):
        gold='Gold Family Plan 31 to 55 Years R 40 000';cremation='Martin`s Cremation Plan For Families'
        self.import_rows([record(product=gold,cover='40000'),record(product=cremation,cover='20000'),record(product='New plan identified by main cover',cover='50000'),record(product='Plan with no main cover supplied',cover='0')])
        records=list(self.store.records('batch1'));self.assertEqual([r['plan'] for r in records],['Plan 2','Historic','Plan 2','Unclassified'])
        batch=self.store.publications()[0];items=catalog(self.store,batch)
        plan=next(i for i in items if i['plan']=='Plan 2');out=self.root/plan['name']
        with patch('central_calculations.fees',side_effect=AssertionError('Export must never recalculate')):metadata=generate(self.store,batch,plan,out)
        self.assertEqual(metadata['rows'],2)
        w=load_workbook(out,data_only=True);self.assertEqual(w['Sheet1'].max_column,19);self.assertEqual(w['Sheet1']['O1'].value,'MPIA');self.assertEqual(w['Sheet1']['F2'].data_type,'s');self.assertEqual(w['Sheet1']['G2'].number_format,'dd/mm/yyyy');w.close()
        standard=next(i for i in items if i['kind']=='policy' and i['destination']=='AUL' and not i['plan'])
        with self.assertRaisesRegex(ValueError,'unclassified'):generate(self.store,batch,standard,self.root/standard['name'])
        rec=next(i for i in items if i['kind']=='reconciliation');generate(self.store,batch,rec,self.root/rec['name'])
    def test_paid_headers_not_positions_counts_once_and_month_scope(self):
        path=self.root/'claims.xlsx';w=Workbook();s=w.active
        s.append(['Final Amount','Product','Date Of Claim Paid','Participating Group','Policy Number'])
        s.append([100,'Plan A','2026-09-10','ERMELO','same']);s.append([100,'Plan A','2026-09-10','ERMELO','same']);s.append([0,'Plan A','2026-09-10','ERMELO','same'])
        w.save(path);records=list(read_transactions(path,'Final Amount'));self.assertEqual(len(records),2)
        with sqlite3.connect(':memory:') as stage:
            stage.execute('CREATE TABLE claims(n INTEGER PRIMARY KEY,month TEXT,franchise TEXT,product TEXT,amount TEXT,source TEXT,issues TEXT)')
            for r in records:stage.execute('INSERT INTO claims VALUES (?,?,?,?,?,?,?)',(r['n'],r['month'],r['franchise'],r['product'],r['amount'],json.dumps(r['source']),'[]'))
            with self.assertRaises(ValueError):publish_claims(self.store,stage,'claims.xlsx',1,'Final Amount',['2026-08'])
            publish_claims(self.store,stage,'claims.xlsx',1,'Final Amount',['2026-09'])
        self.assertEqual(sum(Decimal(r['amount']) for r in self.store.claims()),200)
        self.assertEqual(len(list(self.store.claims())),2)
        with self.assertRaisesRegex(ValueError,'heading'):list(read_transactions(path,''))
    def test_rank_population_ties_and_prior_year_movement(self):
        ranks=retail_rank({'B':100,'A':100,'C':50},{'B':200,'A':100})
        self.assertEqual([r['franchise'] for r in ranks],['A','B','C'])
        self.assertEqual([r['change'] for r in ranks],[1,-1,None])
    def test_insurer_calculation_keeps_bank_claims_and_fees_separate(self):
        self.import_rows([record(franchise='ERMELO',risk='200',retail='300')])
        batch=self.store.publications()[0];item=next(i for i in catalog(self.store,batch) if i['kind']=='payover')
        path=self.root/'payover.xlsx';generate(self.store,batch,item,path)
        book=load_workbook(path,data_only=True)
        self.assertEqual(book['Monthly Pay-over']['F2'].value,15)
        self.assertEqual(book['Insurer Calculation']['L2'].value,'Unconfirmed paid claims');book.close()
        with sqlite3.connect(':memory:') as stage:
            stage.execute('CREATE TABLE claims(n INTEGER PRIMARY KEY,month TEXT,franchise TEXT,product TEXT,amount TEXT,source TEXT,issues TEXT)')
            stage.execute('INSERT INTO claims VALUES (?,?,?,?,?,?,?)',(2,'2026-09','ERMELO','Plan A','50',json.dumps({'Net amount paid via bank':50,'Amount net off against claim':20,'Final Amount':70}),'[]'))
            publish_claims(self.store,stage,'claims.xlsx',1,'Net amount paid via bank',['2026-09'])
        generate(self.store,batch,item,path);book=load_workbook(path,data_only=True)
        sheet=book['Insurer Calculation']
        self.assertEqual(sheet['J2'].value,50)
        self.assertEqual(sheet['K2'].value,95)
        self.assertAlmostEqual(sheet['L2'].value,sheet['I2'].value-95)
        book.close()
    def test_durable_job_checksums_failures_do_not_stop_other_months(self):
        app=Flask(__name__);audit=ImportAudit(self.root/'audit.sqlite');queue=JobQueue(app,self.store,self.root,audit)
        class Upload:
            def __init__(self,rows):
                out=io.StringIO();csv.writer(out).writerows(rows);self.stream=io.BytesIO(out.getvalue().encode())
        with patch.object(queue,'kick'):
            good=queue.create(1,'direct',{'filename':'PolicyData_20260901_to_20260930.csv'},Upload([record()]))
            bad=queue.create(1,'direct',{'filename':'PolicyData_20261001_to_20261031.csv'},Upload([record(policy='')]))
            later=queue.create(1,'direct',{'filename':'PolicyData_20261101_to_20261130.csv'},Upload([record(risk='200')]))
            queue.run()
            self.assertEqual(queue.status(good,1)['state'],'completed');self.assertEqual(queue.status(bad,1)['state'],'failed');self.assertEqual(queue.status(later,1)['state'],'completed');self.assertIsNone(queue.status(good,2))
            self.assertEqual({b['month'] for b in self.store.publications()},{'2026-09','2026-11'})
            self.assertEqual(len(audit.rows(1)),3)

if __name__=='__main__':unittest.main()
